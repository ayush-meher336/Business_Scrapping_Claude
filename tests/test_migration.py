import datetime as dt

import openpyxl
import pytest

from leadforge.core.ids import new_ulid
from leadforge.migration import workbook as wbm
from leadforge.migration.importer import is_role_based, migrate, split_emails, to_text
from leadforge.migration.verify import verify
from leadforge.storage import db


def _sheet(wb, name, rows):
    ws = wb.create_sheet(name)
    for r in rows:
        ws.append(r)


def _lead(lead_id, company, website, suburb, email="info@a.com.au"):
    row = [None] * len(wbm.LEADS_HEADERS)
    fields = [f for _, f in wbm.LEADS_HEADERS]
    values = {
        "lead_id": lead_id, "region": "Oceania", "country": "Australia", "city": "Sydney",
        "suburb_area": suburb, "company_name": company, "website": website, "postcode": 2164,
        "general_email": email, "phone": "02 9756 2636", "source_url": website,
        "source_type": "Company website - Contact page", "email_verified": "Not checked",
        "date_collected": dt.datetime(2026, 9, 24), "legacy_ok_to_email": "No",
    }
    for k, v in values.items():
        row[fields.index(k)] = v
    return row


@pytest.fixture
def workbook(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _sheet(wb, "How To Use", [["ScrapTrade tracker"], ["LEGEND"]])
    _sheet(wb, "Leads", [
        [h for h, _ in wbm.LEADS_HEADERS],
        _lead(1, "W & K Scrap Metal", "https://wk.com.au/", "Wetherill Park", "w.k@bigpond.com; sales@wk.com.au"),
        _lead(2, "Safari Copper", "https://safari.com.au/", "Smithfield"),
        _lead(None, None, None, None, None),  # formula-only row: no company
    ])
    _sheet(wb, "Countries", [wbm.COUNTRIES_HEADERS,
                             [1, "Oceania", "Australia", "Home market", "English", "scrap metal, copper buyer", "Spam Act", None,
                              2, 1, 2, 80, 2, 0.025, 0],
                             ["Footnote text, not a country"]])
    _sheet(wb, "Cities", [wbm.CITIES_HEADERS,
                          [1, 1, "A", "Oceania", "Australia", "NSW", "Sydney", "X", "Metro / Port", "Yards",
                           "Wetherill Park, Smithfield", "Largest metro", "scrap metal Sydney", 2, 80, 2, 0.025,
                           0, 2, "Done", "Claude", "Done 06-Oct"]])
    _sheet(wb, "Search Zones", [wbm.ZONES_HEADERS,
                                [1, 1, "Australia", "Sydney", "Wetherill Park", "X", "A", "scrap metal Wetherill Park",
                                 1, "Done", "Claude", dt.datetime(2026, 9, 24), "Pass 1"],
                                [2, 1, "Australia", "Sydney", "Smithfield", "X", "A", "scrap metal Smithfield",
                                 1, "Done", "Claude", dt.datetime(2026, 9, 24), None]])
    _sheet(wb, "Summary", [["Pipeline summary"], ["Total leads", 2]])
    _sheet(wb, "Sources", [wbm.SOURCES_HEADERS, ["Google Maps", "Global", "Yards", "Search", "Name, phone"]])
    _sheet(wb, "Compliance", [["Cold-email rules"], [], wbm.COMPLIANCE_HEADERS,
                              ["Australia", "Spam Act 2003", "Inferred consent", "Record Source URL"]])
    _sheet(wb, "Research Log", [wbm.LOG_HEADERS,
                                [dt.datetime(2026, 9, 24), "Australia", "Sydney", "Wetherill Park",
                                 "scrap metal Wetherill Park", "Google Maps", 8, 2, "ok"]])
    _sheet(wb, "Lists", [["Industry Type", "Material Focus", "Size code", "Meaning", "Target leads per city"],
                         ["Scrap Yard / Metal Recycler", "Mixed metals", "X", "Mega metro", 80],
                         ["Scrap Buyer / Trader", "Batteries", "M", "Major city", 40],
                         [None, None, "Blue targets are suggestions", None, None]])
    path = tmp_path / "tracker.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def conn(tmp_path):
    c = db.connect(tmp_path / "t.db")
    db.init_schema(c)
    yield c
    c.close()


def test_migrate_round_trip(workbook, conn):
    stats = migrate(conn, workbook)
    assert stats["leads"] == 2
    assert stats["lead_rows_without_company"] == 1
    assert stats["countries"] == 1
    assert stats["zones"] == 2 and stats["zones_without_city"] == 0
    assert stats["leads_linked_to_zone"] == 2
    res = verify(conn, workbook)
    assert res.ok, res.mismatches


def test_lead_ids_frozen_and_keys_unique(workbook, conn):
    migrate(conn, workbook)
    rows = conn.execute("SELECT lead_id, lead_key, postcode FROM business ORDER BY lead_id").fetchall()
    assert [r["lead_id"] for r in rows] == [1, 2]
    assert len({r["lead_key"] for r in rows}) == 2
    assert rows[0]["postcode"] == "2164"


def test_contacts_split_and_classified(workbook, conn):
    migrate(conn, workbook)
    emails = {r["value"]: r["role_based"] for r in conn.execute("SELECT * FROM contact WHERE kind='email'")}
    assert emails == {"w.k@bigpond.com": 0, "sales@wk.com.au": 1, "info@a.com.au": 1}


def test_enums_and_city_sizes(workbook, conn):
    migrate(conn, workbook)
    sizes = conn.execute("SELECT code, target_leads FROM city_size ORDER BY code").fetchall()
    assert [tuple(r) for r in sizes] == [("M", 40), ("X", 80)]
    assert conn.execute("SELECT COUNT(*) FROM enum_value WHERE list_name='Industry Type'").fetchone()[0] == 2


def test_header_change_is_rejected(workbook, conn):
    wb = openpyxl.load_workbook(workbook)
    wb["Leads"]["G1"] = "Business Name"
    wb.save(workbook)
    with pytest.raises(wbm.WorkbookFormatError, match="Company Name"):
        migrate(conn, workbook)


def test_verify_detects_tampering(workbook, conn):
    migrate(conn, workbook)
    conn.execute("UPDATE business SET phone='000' WHERE lead_id=1")
    assert not verify(conn, workbook).ok


def test_helpers():
    assert split_emails("A@x.com, a@x.com; b@y.org.") == ["a@x.com", "b@y.org"]
    assert is_role_based("enquiries@x.com") and not is_role_based("john@x.com")
    assert to_text(dt.datetime(2026, 9, 24)) == "2026-09-24"
    assert to_text(2164.0) == "2164"
    a, b = new_ulid(), new_ulid()
    assert len(a) == 26 and a != b
    assert new_ulid(1) < new_ulid(2)

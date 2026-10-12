"""Read the legacy ScrapTrade workbook: cached cell values, plus the formula each column used."""
from __future__ import annotations

import datetime as dt
import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl

# Expected headers per tab, in sheet order. A mismatch stops the migration rather than
# silently mapping a column to the wrong field.
LEADS_HEADERS = [
    ("Lead ID", "lead_id"),
    ("Region", "region"),
    ("Country", "country"),
    ("State / Province", "state_province"),
    ("City", "city"),
    ("Suburb / Area", "suburb_area"),
    ("Company Name", "company_name"),
    ("Website", "website"),
    ("Located in zone?", "located_in_zone"),
    ("Street Address", "street_address"),
    ("Postcode", "postcode"),
    ("Business Reg. No. (ABN)", "business_reg_no"),
    ("Registered Name", "registered_name"),
    ("Registration Status", "registration_status"),
    ("Industry Type", "industry_type"),
    ("Material Focus", "material_focus"),
    ("Company Size", "company_size"),
    ("General Email", "general_email"),
    ("Decision Maker Name", "decision_maker_name"),
    ("Decision Maker Role", "decision_maker_role"),
    ("Decision Maker Email", "decision_maker_email"),
    ("Email Type", "email_type"),
    ("Phone", "phone"),
    ("LinkedIn Company Page", "linkedin_company_page"),
    ("Source URL", "source_url"),
    ("Source Type", "source_type"),
    ("Found via (keyword)", "found_via_keyword"),
    ("Opt-out notice on site?", "opt_out_notice"),
    ("Email Verified", "email_verified"),
    ("Priority", "priority"),
    ("Outreach Status", "outreach_status"),
    ("Last Contact Date", "last_contact_date"),
    ("Date Collected", "date_collected"),
    ("Collected By", "collected_by"),
    ("Notes", "notes"),
    ("Duplicate?", "legacy_duplicate_flag"),
    ("OK to Email?", "legacy_ok_to_email"),
    ("Website Status", "website_status"),
    ("Last Verified", "last_verified"),
]

COUNTRIES_HEADERS = [
    "Tier", "Region", "Country", "Trade role", "Language(s)",
    "Local search keywords (use in Google Maps)",
    "Cold-email rule (summary - verify before sending)", "Notes / restrictions",
    "Cities listed", "Priority A cities", "Search zones", "Target leads",
    "Leads collected", "% of target", "OK-to-email leads",
]

CITIES_HEADERS = [
    "Loc ID", "Tier", "City Priority", "Region", "Country", "State / Province", "City / Hub",
    "Size", "City type", "Who to target here", "Industrial areas / suburbs to search",
    "Scrap & industry profile", "Suggested Google Maps searches", "Search zones",
    "Target leads", "Leads collected", "% of target", "OK-to-email leads", "Zones done",
    "Status", "Assigned to", "Notes",
]

ZONES_HEADERS = [
    "Zone ID", "Tier", "Country", "City / Hub", "Suburb / Area", "City size", "City Priority",
    "Google Maps search", "Leads collected", "Status", "Assigned to", "Date done", "Notes",
]

LOG_HEADERS = [
    "Date", "Country", "City", "Suburb / Zone", "Keyword / query", "Source",
    "Results reviewed", "New companies added", "Notes",
]

SOURCES_HEADERS = ["Source", "Scope", "Best for (lead type)", "Where / how", "What you typically get"]

COMPLIANCE_HEADERS = [
    "Market", "Law", "B2B cold email to a published business address?", "What you must do",
]

EXPECTED_TABS = [
    "How To Use", "Leads", "Countries", "Cities", "Search Zones", "Summary",
    "Sources", "Compliance", "Research Log", "Lists",
]


class WorkbookFormatError(ValueError):
    pass


@dataclass
class Tab:
    name: str
    width: int
    rows: list[tuple[int, list[Any]]] = field(default_factory=list)  # (sheet row, values)

    def header_row(self, first_header: str) -> tuple[int, list[Any]]:
        for row_num, values in self.rows:
            if values and values[0] == first_header:
                return row_num, values
        raise WorkbookFormatError(f"{self.name}: header row starting {first_header!r} not found")

    def data_after(self, header_row_num: int) -> list[tuple[int, list[Any]]]:
        return [(n, v) for n, v in self.rows if n > header_row_num]


def is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def to_json_value(value: Any) -> Any:
    if isinstance(value, dt.datetime):
        return value.isoformat()
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    return value


def row_json(values: list[Any]) -> str:
    return json.dumps([to_json_value(v) for v in values], ensure_ascii=False)


def read_values(path: Path) -> dict[str, Tab]:
    """Every non-blank row of every tab, using the values Excel last calculated."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        missing = [t for t in EXPECTED_TABS if t not in wb.sheetnames]
        if missing:
            raise WorkbookFormatError(f"workbook is missing tabs: {missing}")
        tabs: dict[str, Tab] = {}
        for ws in wb.worksheets:
            width = ws.max_column or 0
            tab = Tab(ws.title, width)
            for row_num, row in enumerate(ws.iter_rows(values_only=True), start=1):
                values = list(row) + [None] * (width - len(row))
                if all(is_blank(v) for v in values):
                    continue
                tab.rows.append((row_num, values))
            tabs[ws.title] = tab
        return tabs
    finally:
        wb.close()


def read_formulas(path: Path, tabs: dict[str, Tab]) -> dict[str, dict[str, str]]:
    """The formula in each column of the first data row (row 2), keyed by header."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
    try:
        result: dict[str, dict[str, str]] = {}
        for name in ("Leads", "Countries", "Cities", "Search Zones"):
            ws = wb[name]
            rows = list(ws.iter_rows(min_row=1, max_row=2, values_only=True))
            if len(rows) < 2:
                continue
            headers, first = rows
            result[name] = {
                str(h): v for h, v in zip(headers, first)
                if h is not None and isinstance(v, str) and v.startswith("=")
            }
        return result
    finally:
        wb.close()


def check_headers(tab: Tab, header_row: list[Any], expected: list[str]) -> None:
    actual = [h for h in header_row[: len(expected)]]
    if actual != expected:
        diffs = [
            f"col {i + 1}: expected {e!r}, found {a!r}"
            for i, (e, a) in enumerate(zip(expected, actual)) if e != a
        ]
        raise WorkbookFormatError(f"{tab.name}: unexpected headers: " + "; ".join(diffs))
